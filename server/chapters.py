import hashlib
import re
from pathlib import Path
from threading import Lock

from fastapi import HTTPException
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import Chapter, ChapterDetail, Slide


def natural_key(path: Path):
    return [int(part) if part.isdigit() else part.lower()
            for part in re.split(r"(\d+)", path.name)]


def shape_text(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shape_text(shape.shapes)
        elif shape.has_text_frame:
            if shape.text.strip():
                yield shape.text.strip()
        elif shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def meaningful_text(text: str) -> str:
    """Remove footer-only content when deciding whether a slide is readable."""
    return "\n".join(
        line for line in text.splitlines()
        if line.strip() and not re.search(r"copyright|[©®]|^\s*\d+\s*$", line, re.I)
    ).strip()


class ChapterStore:
    def __init__(self, book_dir: Path, slides_dir: Path):
        self.book_dir = book_dir.resolve()
        self.slides_dir = slides_dir.resolve()
        self._cache = {}
        self._lock = Lock()

    def files(self):
        found = {}
        for path in sorted(self.book_dir.glob("*.pptx"), key=natural_key):
            if path.name.startswith("~$") or not path.resolve().is_relative_to(self.book_dir):
                continue
            match = re.search(r"chapter[\s_-]*(\d+)", path.stem, re.I)
            chapter_id = ("chapter-" + str(int(match[1])) if match else
                          "chapter-" + hashlib.sha256(path.name.encode()).hexdigest()[:12])
            if chapter_id in found:
                raise HTTPException(500, "Duplicate chapter number in book directory")
            found[chapter_id] = path
        return found

    def file(self, chapter_id: str):
        path = self.files().get(chapter_id)
        if path is None:
            raise HTTPException(404, "Chapter not found")
        return path

    def image(self, chapter_id: str, number: int):
        path = (self.slides_dir / chapter_id / f"{number}.png").resolve()
        if path.is_relative_to(self.slides_dir) and path.is_file():
            return path
        return None

    def read(self, chapter_id: str) -> ChapterDetail:
        path = self.file(chapter_id)
        stat = path.stat()
        stamp = (stat.st_mtime_ns, stat.st_size)
        with self._lock:
            cached = self._cache.get(chapter_id)
            if cached is None or cached[0] != stamp:
                try:
                    presentation = Presentation(path)
                    slides = []
                    for number, slide in enumerate(presentation.slides, 1):
                        text = "\n".join(shape_text(slide.shapes))
                        heading = slide.shapes.title
                        title = heading.text.strip() if heading is not None else ""
                        if not title:
                            lines = meaningful_text(text).splitlines()
                            title = lines[0][:300] if lines else f"Slide {number}"
                        notes = ""
                        if slide.has_notes_slide:
                            frame = slide.notes_slide.notes_text_frame
                            notes = frame.text.strip() if frame is not None else ""
                        slides.append(Slide(number=number, title=title[:500],
                                            text=text, notes=notes))
                    chapter = ChapterDetail(
                        id=chapter_id, title=slides[0].title if slides else path.stem,
                        filename=path.name, slideCount=len(slides), slides=slides,
                    )
                except Exception as exc:
                    raise HTTPException(500, "Unable to parse chapter presentation") from exc
                self._cache[chapter_id] = (stamp, chapter)
            result = self._cache[chapter_id][1].model_copy(deep=True)
        for slide in result.slides:
            if self.image(chapter_id, slide.number):
                slide.imageUrl = f"/api/chapters/{chapter_id}/slides/{slide.number}/image"
        return result

    def list(self):
        return [Chapter.model_validate(self.read(chapter_id).model_dump(exclude={"slides"}))
                for chapter_id in self.files()]
