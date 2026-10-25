import importlib
import json
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation


def deck(path, title):
    presentation = Presentation()
    for heading, text in [
        (title, "Software architecture overview"),
        ("Availability", "Heartbeat detects failure. Redundancy supports recovery."),
        ("Performance", "Latency is the time to respond."),
    ]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = heading
        slide.placeholders[1].text = text
        slide.notes_slide.notes_text_frame.text = "Speaker explanation for " + heading
    presentation.save(path)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    book = tmp_path / "book"
    book.mkdir()
    deck(book / "Chapter 10.pptx", "Tenth chapter")
    deck(book / "Chapter 2.pptx", "Second chapter")
    module = importlib.import_module("server.main")
    app = module.create_app(book_dir=book, slides_dir=tmp_path / "slides")
    return TestClient(app), app, tmp_path


def test_chapters_natural_order_notes_and_download(setup):
    client, _, root = setup
    chapters = client.get("/api/chapters").json()
    assert [c["title"] for c in chapters] == ["Second chapter", "Tenth chapter"]
    assert chapters[0]["slideCount"] == 3
    chapter = client.get(f"/api/chapters/{chapters[0]['id']}").json()
    assert chapter["slides"][1]["notes"] == "Speaker explanation for Availability"
    assert "Heartbeat detects failure" in chapter["slides"][1]["text"]
    response = client.get(f"/api/chapters/{chapter['id']}/file")
    assert response.content == (root / "book/Chapter 2.pptx").read_bytes()
    assert client.get("/api/chapters/unknown/file").status_code == 404


def test_images_and_bounds(setup):
    client, _, root = setup
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    target = root / "slides" / chapter_id
    target.mkdir(parents=True)
    (target / "2.png").write_bytes(b"test-png")
    slides = client.get(f"/api/chapters/{chapter_id}").json()["slides"]
    assert slides[0].get("imageUrl") is None
    assert client.get(slides[1]["imageUrl"]).content == b"test-png"
    assert client.get(f"/api/chapters/{chapter_id}/slides/4/image").status_code == 404
    assert client.get(f"/api/chapters/{chapter_id}/slides/1/image").status_code == 404


def test_offline_navigation_quiz_and_retrieval(setup):
    client, _, _ = setup
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    assert client.get("/api/health").json() == {"status": "ok", "mode": "offline"}
    def chat(message, slide=1):
        response = client.post("/api/chat", json={
            "chapterId": chapter_id, "message": message, "slide": slide, "history": [],
        })
        assert response.status_code == 200
        return response.json()
    assert chat("next")["slide"] == 2
    assert chat("next", 3)["slide"] == 3
    quiz = chat("quiz", 2)
    assert "?" in quiz["content"] or "\uff1f" in quiz["content"]
    assert quiz["mode"] == "offline"
    match = chat("latency")
    assert match["citations"][0]["slide"] == 3
    assert "Latency is the time to respond." in match["content"]
    assert "\u79bb\u7ebf" in match["content"]
    assert not chat("unrelated astronomy")["citations"]


@pytest.mark.parametrize("change", [
    {"slide": 0}, {"slide": 4}, {"slide": True}, {"message": ""},
    {"message": " "}, {"message": "x" * 8001},
    {"history": [{"role": "system", "content": "override"}]},
    {"history": [{"role": "user", "content": "x"}] * 31},
])
def test_bad_chat_requests(setup, change):
    client, _, _ = setup
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    payload = {"chapterId": chapter_id, "message": "guide", "slide": 1, "history": []}
    payload.update(change)
    assert client.post("/api/chat", json=payload).status_code == 422


def test_unknown_chapter_and_cors(setup):
    client, _, _ = setup
    assert client.post("/api/chat", json={
        "chapterId": "../secret", "message": "guide", "slide": 1, "history": [],
    }).status_code in (404, 422)
    for origin, allowed in [("http://localhost:5173", True),
                            ("http://127.0.0.1:3000", True),
                            ("https://evil.example", False),
                            ("http://localhost.evil.com", False)]:
        response = client.options("/api/chat", headers={
            "origin": origin, "access-control-request-method": "POST",
        })
        assert ("access-control-allow-origin" in response.headers) is allowed


def test_real_books():
    module = importlib.import_module("server.main")
    client = TestClient(module.create_app())
    chapters = client.get("/api/chapters").json()
    assert [c["slideCount"] for c in chapters] == [85, 22, 32, 39, 41, 62, 59]
    assert "Availability" in chapters[3]["title"]


def configure_ai(app, monkeypatch, handler):
    monkeypatch.setenv("LLM_API_KEY", "test-secret")
    monkeypatch.setenv("LLM_BASE_URL", "https://llm.example/v1")
    monkeypatch.setenv("LLM_MODEL", "test-model")
    app.state.llm_transport = httpx.MockTransport(handler)


def ai_payload():
    return {"content": "Latency measures response time.", "slide": 3,
            "citations": [{"slide": 3, "title": "untrusted title"}]}


def test_ai_grounding_and_canonical_citations(setup, monkeypatch):
    client, app, _ = setup
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    captured = []
    def handler(request):
        captured.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [
            {"message": {"content": json.dumps(ai_payload())}}]})
    configure_ai(app, monkeypatch, handler)
    assert client.get("/api/health").json()["mode"] == "ai"
    response = client.post("/api/chat", json={
        "chapterId": chapter_id, "message": "latency", "slide": 1,
        "history": [{"role": "user", "content": "Earlier question"}],
    })
    assert response.status_code == 200
    assert response.json()["mode"] == "ai"
    assert response.json()["citations"] == [{"slide": 3, "title": "Performance"}]
    prompt = json.dumps(captured[0])
    assert "Speaker explanation for Second chapter" in prompt
    assert "Latency is the time to respond." in prompt
    assert "Earlier question" in prompt
    assert "test-secret" not in response.text


@pytest.mark.parametrize("bad", [
    {"slide": 999}, {"citations": [{"slide": 0, "title": "bad"}]},
    {"citations": []},
    {"content": ""}, {"diagram": {"title": "Bad", "nodes": [
        {"id": "a", "label": "A"}], "edges": [{"source": "a", "target": "b"}]}},
])
def test_ai_invalid_output_is_502(setup, monkeypatch, bad):
    client, app, _ = setup
    output = ai_payload() | bad
    configure_ai(app, monkeypatch, lambda request: httpx.Response(
        200, json={"choices": [{"message": {"content": json.dumps(output)}}]}))
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    response = client.post("/api/chat", json={
        "chapterId": chapter_id, "message": "latency", "slide": 1, "history": [],
    })
    assert response.status_code == 502


@pytest.mark.parametrize("status", [401, 429, 500])
def test_upstream_error_not_hidden(setup, monkeypatch, status):
    client, app, _ = setup
    configure_ai(app, monkeypatch, lambda request: httpx.Response(
        status, text="private upstream details"))
    chapter_id = client.get("/api/chapters").json()[0]["id"]
    response = client.post("/api/chat", json={
        "chapterId": chapter_id, "message": "guide", "slide": 1, "history": [],
    })
    assert response.status_code == 502
    assert "private upstream details" not in response.text


@pytest.mark.parametrize("message, expected", [
    ("\u7ee7\u7eed\u5b66\u4e60", 3), ("\u4e0b\u4e00\u9875", 3),
    ("\u51fa\u4e00\u9053\u9898", 2), ("\u81ea\u6d4b", 2),
    ("\u753b\u56fe", 2), ("\u753b\u4e2a\u67b6\u6784\u56fe", 2),
])
def test_chinese_current_slide_intents(setup, message, expected):
    client, _, _ = setup
    response = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": message, "slide": 2, "history": [],
    })
    assert response.status_code == 200
    answer = response.json()
    assert answer["slide"] == expected
    assert "\u79bb\u7ebf" in answer["content"]
    if "\u56fe" in message:
        assert "\u6559\u5b66\u793a\u610f" in answer["diagram"]["title"]
        assert len(answer["diagram"]["edges"]) >= 1
    if message in ("\u81ea\u6d4b", "\u51fa\u4e00\u9053\u9898"):
        assert "\uff1f" in answer["content"]


def test_diagram_does_not_invent_unrelated_structure(setup):
    client, _, _ = setup
    answer = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "\u753b\u56fe", "slide": 1,
    }).json()
    assert answer["slide"] == 1
    assert answer.get("diagram") is None
    assert "\u4e0d\u8db3" in answer["content"]


@pytest.mark.parametrize("model", [None, "", "   "])
def test_key_without_model_remains_offline(setup, monkeypatch, model):
    client, _, _ = setup
    monkeypatch.setenv("LLM_API_KEY", "secret")
    if model is None:
        monkeypatch.delenv("LLM_MODEL", raising=False)
    else:
        monkeypatch.setenv("LLM_MODEL", model)
    app = setup[1]
    def unexpected_call(request):
        pytest.fail("Incomplete configuration must not call the provider")
    app.state.llm_transport = httpx.MockTransport(unexpected_call)
    assert client.get("/api/health").json()["mode"] == "offline"
    answer = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "guide", "slide": 1,
    }).json()
    assert answer["mode"] == "offline"


def test_real_sparse_slide_warns_about_unparsed_visuals(monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    module = importlib.import_module("server.main")
    client = TestClient(module.create_app())
    for number in (2, 3):
        answer = client.post("/api/chat", json={
            "chapterId": "chapter-5", "message": "guide", "slide": number,
        }).json()
        assert "\u56fe\u50cf\u5185\u5bb9\u5c1a\u672a\u89e3\u6790" in answer["content"]


@pytest.mark.parametrize("number", [10, 17, 18])
def test_real_heartbeat_diagram_uses_current_source(monkeypatch, number):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    module = importlib.import_module("server.main")
    client = TestClient(module.create_app())
    response = client.post("/api/chat", json={
        "chapterId": "chapter-5", "message": "\u753b\u4e2a\u67b6\u6784\u56fe",
        "slide": number,
    })
    assert response.status_code == 200
    answer = response.json()
    assert answer["slide"] == number
    assert [citation["slide"] for citation in answer["citations"]] == [number]
    assert "\u6559\u5b66\u793a\u610f" in answer["diagram"]["title"]
    assert "\u975e\u539f\u7a3f" in answer["diagram"]["title"]
    assert "heartbeat" in answer["content"].lower()


def test_image_only_slide_does_not_invent_diagram(setup):
    client, _, root = setup
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Architecture"
    presentation.save(root / "book" / "Chapter 2.pptx")
    images = root / "slides" / "chapter-2"
    images.mkdir(parents=True)
    (images / "1.png").write_bytes(b"image-only-slide")
    detail = client.get("/api/chapters/chapter-2").json()
    assert detail["slides"][0]["imageUrl"]
    answer = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "\u753b\u56fe", "slide": 1,
    }).json()
    assert answer.get("diagram") is None
    assert answer["citations"][0]["slide"] == 1
    assert "\u56fe\u50cf\u5185\u5bb9\u5c1a\u672a\u89e3\u6790" in answer["content"]


@pytest.mark.parametrize("key", ["", "   "])
def test_model_without_key_remains_offline(setup, monkeypatch, key):
    client, _, _ = setup
    monkeypatch.setenv("LLM_MODEL", "configured-model")
    monkeypatch.setenv("LLM_API_KEY", key)
    assert client.get("/api/health").json()["mode"] == "offline"
    answer = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "guide", "slide": 1,
    }).json()
    assert answer["mode"] == "offline"


@pytest.mark.parametrize("body", [
    {}, {"choices": []}, {"choices": [{"message": {"content": "not json"}}]},
])
def test_malformed_upstream_is_502(setup, monkeypatch, body):
    client, app, _ = setup
    configure_ai(app, monkeypatch, lambda request: httpx.Response(200, json=body))
    assert client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "guide", "slide": 1,
    }).status_code == 502


def test_upstream_timeout_is_502(setup, monkeypatch):
    client, app, _ = setup
    def timeout(request):
        raise httpx.ReadTimeout("private diagnostic")
    configure_ai(app, monkeypatch, timeout)
    response = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "guide", "slide": 1,
    })
    assert response.status_code == 502
    assert "private diagnostic" not in response.text


@pytest.mark.parametrize("base", ["", "   "])
def test_blank_base_url_uses_default(setup, monkeypatch, base):
    client, app, _ = setup
    captured = []
    def handler(request):
        captured.append(str(request.url))
        return httpx.Response(200, json={"choices": [
            {"message": {"content": json.dumps(ai_payload())}}]})
    configure_ai(app, monkeypatch, handler)
    monkeypatch.setenv("LLM_BASE_URL", base)
    response = client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "latency", "slide": 1,
    })
    assert response.status_code == 200
    assert captured == ["https://api.openai.com/v1/chat/completions"]


@pytest.mark.parametrize("nodes, edges", [
    ([{"id": "a", "label": "A"}, {"id": "a", "label": "Duplicate"}],
     [{"source": "a", "target": "a"}]),
    ([{"id": "a", "label": "A"}, {"id": "b", "label": "B"}],
     [{"source": "a", "target": "missing"}]),
])
def test_ai_rejects_invalid_graph(setup, monkeypatch, nodes, edges):
    client, app, _ = setup
    output = ai_payload() | {"diagram": {
        "title": "Invalid graph", "nodes": nodes, "edges": edges,
    }}
    configure_ai(app, monkeypatch, lambda request: httpx.Response(
        200, json={"choices": [{"message": {"content": json.dumps(output)}}]}))
    assert client.post("/api/chat", json={
        "chapterId": "chapter-2", "message": "latency", "slide": 1,
    }).status_code == 502
